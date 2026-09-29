"""PPTX support.

Strategy:
1. If LibreOffice (soffice) is installed, convert to PDF for pixel accurate slide images and parse that.
2. Always read the PPTX with python-pptx to recover speaker notes, real chart data and table data,
   which are better grounding than anything we can infer from pixels.
3. Without LibreOffice, render a clean text-based preview image with Pillow.
"""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

from ..analysis.text_utils import key_phrases, normalize, numbers_in
from ..models import SlideContext, VisualElement
from .pdf_parser import clean_bullet, parse_pdf

log = logging.getLogger("pitchmirror.ingest.pptx")

MAX_UNCOMPRESSED = 400 * 1024 * 1024


def check_zip_safety(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            if len(infos) > 10000:
                raise ValueError("The PPTX contains too many internal files.")
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                raise ValueError("The PPTX is too large once decompressed.")
            names = {i.filename for i in infos}
            if "[Content_Types].xml" not in names or not any(n.startswith("ppt/slides/") for n in names):
                raise ValueError("This file is not a valid PowerPoint presentation.")
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is not a valid PowerPoint presentation.") from exc


def _shape_text(shape) -> list[str]:
    out = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = clean_bullet("".join(r.text for r in p.runs))
            if t:
                out.append(t)
    return out


def _chart_element(shape) -> VisualElement:
    chart = shape.chart
    try:
        ctype = str(chart.chart_type).split(".")[-1].split(" ")[0].replace("_", " ").lower()
    except Exception:
        ctype = "chart"
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = normalize(chart.chart_title.text_frame.text)
    except Exception:
        pass
    summary_parts = []
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories][:12]
        if cats:
            summary_parts.append("Categories: " + ", ".join(cats))
        for s in plot.series:
            vals = [v for v in s.values][:12]
            summary_parts.append(f"Series '{s.name}': " + ", ".join(f"{v:g}" if isinstance(v, (int, float)) else str(v) for v in vals))
    except Exception:
        pass
    desc = f"{ctype.capitalize()} chart" + (f" titled '{title}'" if title else "")
    return VisualElement(kind="chart", description=desc, source="pptx", data_summary=" | ".join(summary_parts)[:600])


def _table_element(shape) -> VisualElement:
    tbl = shape.table
    rows = [[normalize(c.text) for c in r.cells] for r in tbl.rows]
    header = rows[0] if rows else []
    return VisualElement(
        kind="table",
        description=f"Table with {max(len(rows) - 1, 0)} rows; columns: {', '.join(h for h in header if h)}",
        source="pptx",
        data_summary="; ".join(" | ".join(r) for r in rows[:6])[:600],
    )


def _extract_pptx(path: Path) -> list[dict]:
    prs = Presentation(str(path))
    slide_w = prs.slide_width or Emu(12192000)
    slide_h = prs.slide_height or Emu(6858000)
    area = float(slide_w) * float(slide_h)
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = normalize(slide.shapes.title.text)
        except Exception:
            pass
        body: list[str] = []
        visuals: list[VisualElement] = []

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
                    walk(shape.shapes)
                    continue
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    visuals.append(_chart_element(shape))
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    visuals.append(_table_element(shape))
                    continue
                if shape.shape_type == 13:  # picture
                    try:
                        frac = float(shape.width) * float(shape.height) / area
                    except Exception:
                        frac = 0
                    if frac >= 0.04:
                        alt = ""
                        try:
                            alt = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")[0]
                        except Exception:
                            pass
                        desc = f"Image occupying ~{int(frac * 100)}% of the slide"
                        if alt:
                            desc += f" (alt text: '{normalize(alt)[:120]}')"
                        visuals.append(VisualElement(kind="image", description=desc, source="pptx"))
                    continue
                for t in _shape_text(shape):
                    if t != title:
                        body.append(t)

        walk(slide.shapes)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = normalize(slide.notes_slide.notes_text_frame.text)
        except Exception:
            pass
        slides.append({"title": title, "body": body, "visuals": visuals, "notes": notes})
    return slides


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def _render_text_slide(data: dict, out_file: Path, width: int) -> None:
    height = int(width * 9 / 16)
    img = Image.new("RGB", (width, height), (250, 250, 252))
    draw = ImageDraw.Draw(img)
    try:
        f_title = ImageFont.load_default(size=int(width * 0.035))
        f_body = ImageFont.load_default(size=int(width * 0.021))
        f_small = ImageFont.load_default(size=int(width * 0.016))
    except TypeError:  # very old Pillow
        f_title = f_body = f_small = ImageFont.load_default()
    pad = int(width * 0.06)
    y = pad
    for line in _wrap(draw, data["title"] or "(untitled slide)", f_title, width - 2 * pad)[:2]:
        draw.text((pad, y), line, fill=(20, 24, 40), font=f_title)
        y += int(width * 0.045)
    y += int(width * 0.015)
    for b in data["body"][:12]:
        for i, line in enumerate(_wrap(draw, b, f_body, width - 2 * pad - 30)[:3]):
            draw.text((pad + (0 if i else 0), y), ("• " if i == 0 else "   ") + line, fill=(55, 60, 80), font=f_body)
            y += int(width * 0.03)
        if y > height - pad * 2:
            break
    vx = pad
    for v in data["visuals"][:3]:
        box_w = int((width - 2 * pad) / 3) - 12
        by = height - pad - int(height * 0.16)
        draw.rounded_rectangle([vx, by, vx + box_w, height - pad], radius=12, outline=(120, 110, 220), width=2)
        for i, line in enumerate(_wrap(draw, f"[{v.kind}] {v.description}", f_small, box_w - 20)[:3]):
            draw.text((vx + 10, by + 8 + i * int(width * 0.02)), line, fill=(90, 80, 180), font=f_small)
        vx += box_w + 18
    draw.text((width - pad - 260, height - int(pad * 0.6)), "text preview (install LibreOffice for exact render)",
              fill=(160, 160, 175), font=f_small)
    img.save(out_file)


def _convert_with_libreoffice(path: Path, workdir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    try:
        subprocess.run(
            [soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(workdir), str(path)],
            check=True, timeout=120, capture_output=True,
        )
        pdf = workdir / (path.stem + ".pdf")
        return pdf if pdf.exists() else None
    except Exception as exc:
        log.warning("LibreOffice conversion failed, falling back to text preview: %s", exc)
        return None


def parse_pptx(path: Path, out_dir: Path, render_width: int, max_slides: int) -> list[SlideContext]:
    check_zip_safety(path)
    try:
        extracted = _extract_pptx(path)
    except Exception as exc:
        raise ValueError("Could not read this PowerPoint file. Try exporting it to PDF.") from exc
    if not extracted:
        raise ValueError("The presentation has no slides.")
    if len(extracted) > max_slides:
        raise ValueError(f"The deck has {len(extracted)} slides; the limit is {max_slides}.")

    with tempfile.TemporaryDirectory() as td:
        pdf = _convert_with_libreoffice(path, Path(td))
        pdf_slides = parse_pdf(pdf, out_dir, render_width, max_slides) if pdf is not None else None

    slides: list[SlideContext] = []
    for i, d in enumerate(extracted, start=1):
        image_name = f"slide_{i:03d}.png"
        visuals = list(d["visuals"])
        if pdf_slides is not None and i <= len(pdf_slides):
            # exact render from LibreOffice; keep pixel-detected visuals the PPTX XML did not describe
            kinds = {v.kind for v in visuals}
            for v in pdf_slides[i - 1].visual_elements:
                if v.kind in kinds or (v.kind in ("diagram", "chart") and "chart" in kinds):
                    continue
                if v.kind == "image" and "image" in kinds:
                    continue
                visuals.append(v)
        else:
            _render_text_slide(d, out_dir / image_name, render_width)
        slides.append(_context_from_pptx(i, d, visuals, image_name))
    return slides


def _context_from_pptx(i: int, d: dict, visuals: list[VisualElement], image_name: str) -> SlideContext:
    all_text = [d["title"], *d["body"]]
    numbers: list[str] = []
    for t in all_text:
        for n in numbers_in(t):
            if n not in numbers:
                numbers.append(n)
    wc = sum(len(t.split()) for t in all_text)
    return SlideContext(
        slide_number=i,
        title=d["title"],
        text=d["body"][:40],
        notes=d["notes"],
        visual_elements=visuals,
        key_concepts=key_phrases(d["body"] or [d["title"]]),
        important_elements=([l for l in d["body"] if numbers_in(l)][:5]
                            + [f"{v.kind}: {v.description}" for v in visuals])[:10],
        numbers=numbers[:20],
        image_file=image_name,
        word_count=wc,
        is_empty=wc == 0 and not visuals,
    )
